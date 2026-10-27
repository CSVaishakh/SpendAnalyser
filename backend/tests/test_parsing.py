import datetime as dt
import io
from decimal import Decimal

import pytest

from app.parsing.normalize import parse_amount, parse_date
from app.parsing.statement import StatementError, parse_statement

HEADER = "Sl. No.,Tran Date,Particulars,Value Date,Tran Type,Cheque Details,Withdrawal,Deposit,Balance Amount\n"


def test_sample_statement_parses_and_reconciles(sample_bytes):
    p = parse_statement("s.csv", sample_bytes)
    assert len(p.transactions) == 357
    assert p.balance_mismatches == 0
    assert not p.reversed_order
    assert p.period == (dt.date(2026, 4, 1), dt.date(2026, 9, 30))
    # Hard-wrapped narration is re-joined without a stray space.
    assert any("reliancefresh@axisbank/UTIB" in t.particulars for t in p.transactions)
    assert not any("@ax isbank" in t.particulars for t in p.transactions)


def test_newest_first_statement_is_reordered():
    csv = HEADER + (
        "3,03-04-2026,UPI/3/C,03-04-2026,TFR,,100.00,,700.00\n"
        "2,02-04-2026,UPI/2/B,02-04-2026,TFR,,200.00,,800.00\n"
        "1,01-04-2026,NEFT-X-SALARY,01-04-2026,TFR,,,1000.00,1000.00\n"
    )
    p = parse_statement("s.csv", csv.encode())
    assert p.reversed_order
    assert [t.date.day for t in p.transactions] == [1, 2, 3]
    assert p.balance_mismatches == 0


def test_preamble_footer_and_repeated_header_are_ignored():
    csv = ("Account Statement\nName: Someone\n\n" + HEADER +
           "1,01/04/2026,UPI/1/SHOP,01/04/2026,TFR,,50.00,,950.00\n" + HEADER +
           "2,02/04/2026,UPI/2/SHOP,02/04/2026,TFR,,50.00,,900.00\n"
           ",,Closing balance,,,,,,900.00\n,,Total,,,,100.00,0.00,\n")
    p = parse_statement("s.csv", csv.encode())
    assert len(p.transactions) == 2
    assert p.transactions[-1].particulars == "UPI/2/SHOP"


def test_single_amount_column_with_drcr():
    csv = ("Date,Narration,Amount,Dr/Cr,Balance\n"
           "01-Apr-2026,SALARY ACME,50000,CR,50000\n"
           "02-Apr-2026,UPI/SWIGGY,450.50,DR,49549.50\n")
    t = parse_statement("s.csv", csv.encode()).transactions
    assert t[0].deposit == Decimal("50000") and t[0].withdrawal == 0
    assert t[1].withdrawal == Decimal("450.50")


def test_month_first_dates_detected():
    csv = ("Date,Description,Debit,Credit,Balance\n"
           "04/13/2026,A,10,,90\n04/14/2026,B,10,,80\n")
    t = parse_statement("s.csv", csv.encode()).transactions
    assert t[0].date == dt.date(2026, 4, 13)


def test_xlsx_statement():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Statement of account"])
    ws.append(HEADER.strip().split(","))
    ws.append([1, dt.datetime(2026, 4, 1), "UPI/1/ZOMATO/zomato@paytm", dt.datetime(2026, 4, 1), "TFR", None, 250.5, None, 749.5])
    ws.append([2, dt.datetime(2026, 4, 2), "NEFT-ACME-SALARY", dt.datetime(2026, 4, 2), "TFR", None, None, 1000, 1749.5])
    buf = io.BytesIO()
    wb.save(buf)
    p = parse_statement("s.xlsx", buf.getvalue())
    assert [t.date for t in p.transactions] == [dt.date(2026, 4, 1), dt.date(2026, 4, 2)]
    assert p.transactions[0].withdrawal == Decimal("250.5")
    assert p.balance_mismatches == 0


def test_html_disguised_as_xls():
    html = ("<html><body><table><tr><td>Statement</td></tr>"
            "<tr><th>Tran Date</th><th>Particulars</th><th>Withdrawal</th><th>Deposit</th><th>Balance</th></tr>"
            "<tr><td>01-04-2026</td><td>UPI/123/UBER</td><td>1,200.00</td><td></td><td>8,800.00</td></tr>"
            "</table></body></html>")
    p = parse_statement("s.xls", html.encode())
    assert p.transactions[0].withdrawal == Decimal("1200.00")


def test_identical_rows_in_one_file_are_kept():
    csv = ("Date,Narration,Withdrawal,Deposit\n"
           "01-04-2026,UPI/TEA STALL,10,\n01-04-2026,UPI/TEA STALL,10,\n")
    p = parse_statement("s.csv", csv.encode())
    assert len({t.fingerprint for t in p.transactions}) == 2


@pytest.mark.parametrize("content,msg", [
    (b"%PDF-1.7 ...", "PDF"),
    (b"a,b,c\n1,2,3\n", "header"),
    (b"", "empty"),
])
def test_bad_files(content, msg):
    with pytest.raises(StatementError, match=msg):
        parse_statement("x", content)


@pytest.mark.parametrize("raw,expected", [
    ("1,23,456.78", Decimal("123456.78")), ("₹ 500", Decimal("500")), ("250.00 Dr", Decimal("-250.00")),
    ("(75.5)", Decimal("-75.5")), ("1,000.00 Cr", Decimal("1000.00")), ("", None), ("-", None),
])
def test_parse_amount(raw, expected):
    assert parse_amount(raw) == expected


@pytest.mark.parametrize("raw,expected", [
    ("01-04-2026", dt.date(2026, 4, 1)), ("01/04/26", dt.date(2026, 4, 1)), ("01-Apr-2026", dt.date(2026, 4, 1)),
    ("1 Sept 2026", dt.date(2026, 9, 1)), ("2026-04-01", dt.date(2026, 4, 1)), ("01/04/2026 10:15:00", dt.date(2026, 4, 1)),
    ("31-02-2026", None), ("Opening", None),
])
def test_parse_date(raw, expected):
    assert parse_date(raw) == expected


def test_user_excel_layout_date_details_ref_debit_credit_balance():
    """Layout from a real user statement: Date | Details | Ref No/Cheque No | Debit | Credit | Balance."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Account Name : VAISHAKH S"])
    ws.append(["Account Number : XXXXXXXX4417"])
    ws.append([])
    ws.append(["Date", "Details", "Ref No/Cheque No", "Debit", "Credit", "Balance"])
    ws.append(["01 Apr 2026", "UPI/412345678901/SWIGGY/swiggy@icici\nPayment", "412345678901", "250.00", "", "9,750.00 Cr"])
    ws.append([dt.datetime(2026, 4, 2), "NEFT-ACME TECHNOLOGIES-SALARY", "N123456", None, 50000, "59,750.00 Cr"])
    ws.append(["03/04/2026", "TO CLG/RAMESH", "000123", "1,000.00", "", "58,750.00 Cr"])
    ws.append(["", "Total", "", "1,250.00", "50,000.00", ""])
    buf = io.BytesIO()
    wb.save(buf)
    p = parse_statement("statement.xlsx", buf.getvalue())
    t = p.transactions
    assert [x.date for x in t] == [dt.date(2026, 4, 1), dt.date(2026, 4, 2), dt.date(2026, 4, 3)]
    assert t[0].particulars == "UPI/412345678901/SWIGGY/swiggy@icici Payment"
    assert t[0].cheque == "412345678901"
    assert (t[0].withdrawal, t[1].deposit, t[2].balance) == (Decimal("250.00"), Decimal("50000"), Decimal("58750.00"))
    assert p.balance_mismatches == 0


def _encrypted_xlsx(password: str) -> bytes:
    from msoffcrypto.format.ooxml import OOXMLFile
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Account Number : XXXXXXXX7788"])
    ws.append(["Date", "Details", "Ref No/Cheque No", "Debit", "Credit", "Balance"])
    ws.append(["01-04-2026", "UPI/1/SWIGGY/swiggy@icici", "1", "200.00", "", "800.00"])
    plain, enc = io.BytesIO(), io.BytesIO()
    wb.save(plain)
    plain.seek(0)
    f = OOXMLFile(plain)
    f.encrypt(password, enc)
    return enc.getvalue()


def test_password_protected_xlsx():
    from app.parsing.statement import StatementLocked

    data = _encrypted_xlsx("01011990")
    with pytest.raises(StatementLocked) as e:
        parse_statement("Email_Statement.xlsx", data)
    assert not e.value.wrong_password and "password-protected" in str(e.value)
    with pytest.raises(StatementLocked) as e:
        parse_statement("Email_Statement.xlsx", data, "wrong")
    assert e.value.wrong_password
    p = parse_statement("Email_Statement.xlsx", data, "01011990")
    assert p.transactions[0].withdrawal == Decimal("200.00") and p.account.last4 == "7788"
