"""Turn an uploaded statement file (CSV, XLSX, XLS or HTML-as-XLS) into a grid of strings."""

from __future__ import annotations

import csv
import datetime as dt
import io
from html.parser import HTMLParser

Grid = list[list[str]]

XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class UnsupportedFile(ValueError):
    pass


class PasswordRequired(UnsupportedFile):
    """The workbook is encrypted; `wrong` is True when a password was given but didn't open it."""

    def __init__(self, wrong: bool = False):
        self.wrong = wrong
        super().__init__(
            "That password didn't open the file. Check the password your bank mentioned in the statement email."
            if wrong else
            "This statement is password-protected. Enter the password your bank gave you (often based on your date of "
            "birth, PAN, customer ID or mobile number) to open it."
        )


def _is_encrypted_office(data: bytes) -> bool:
    if not data.startswith(XLS_MAGIC):
        return False
    import olefile

    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except OSError:
        return False
    try:
        # Encrypted .xlsx: an OLE container holding EncryptionInfo + EncryptedPackage.
        if ole.exists("EncryptionInfo") and ole.exists("EncryptedPackage"):
            return True
        # Encrypted legacy .xls: FILEPASS record inside the Workbook stream.
        if ole.exists("Workbook"):
            import msoffcrypto

            return msoffcrypto.OfficeFile(io.BytesIO(data)).is_encrypted()
        return False
    finally:
        ole.close()


def _decrypt(data: bytes, password: str) -> bytes:
    import msoffcrypto
    from msoffcrypto.exceptions import DecryptionError, InvalidKeyError

    f = msoffcrypto.OfficeFile(io.BytesIO(data))
    try:
        f.load_key(password=password, verify_password=True)
        out = io.BytesIO()
        f.decrypt(out)
    except (InvalidKeyError, DecryptionError) as e:
        raise PasswordRequired(wrong=True) from e
    except Exception as e:  # msoffcrypto signals a bad key in several ways depending on the format
        if "key" in str(e).lower() or "password" in str(e).lower():
            raise PasswordRequired(wrong=True) from e
        raise
    return out.getvalue()


def read_grid(filename: str, data: bytes, password: str | None = None) -> Grid:
    if not data.strip():
        raise UnsupportedFile("The file is empty.")
    if _is_encrypted_office(data):
        if not password:
            raise PasswordRequired()
        data = _decrypt(data, password)
    head = data[:4096].lstrip().lower()
    if data.startswith(XLSX_MAGIC):
        grid = _read_xlsx(data)
    elif data.startswith(XLS_MAGIC):
        grid = _read_xls(data)
    elif head.startswith(b"%pdf"):
        raise UnsupportedFile("PDF statements aren't supported. Download the statement as CSV or Excel from your bank.")
    elif head.startswith(b"<") and b"<table" in data[:500_000].lower():
        # Many Indian banks serve an HTML table with an .xls extension.
        grid = _read_html(_decode(data))
    else:
        grid = _read_csv(_decode(data))
    return [[_clean(c) for c in row] for row in grid]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value) if not value.is_integer() else str(int(value))
    return " ".join(str(value).replace("\xa0", " ").split())


def _read_csv(text: str) -> Grid:
    sample = text[:8000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = max(",;\t|", key=sample.count)
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_xlsx(data: bytes) -> Grid:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = _pick_sheet(wb.worksheets, lambda s: list(s.iter_rows(values_only=True)))
        return [list(r) for r in ws]
    finally:
        wb.close()


def _read_xls(data: bytes) -> Grid:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)

    def rows(sheet):
        out = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            out.append(row)
        return out

    return _pick_sheet(book.sheets(), rows)


def _pick_sheet(sheets, to_rows) -> list:
    """Statements sometimes ship with a cover sheet; use the sheet with the most non-empty rows."""
    best: list = []
    for sheet in sheets:
        rows = to_rows(sheet)
        if sum(1 for r in rows if any(v not in (None, "") for v in r)) > sum(1 for r in best if any(v not in (None, "") for v in r)):
            best = rows
    return best


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: Grid = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append("".join(self._cell))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _read_html(text: str) -> Grid:
    p = _TableParser()
    p.feed(text)
    return p.rows
